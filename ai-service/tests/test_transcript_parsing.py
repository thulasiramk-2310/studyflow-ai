"""Exported meeting transcripts become plain `Speaker: text` lines."""

from __future__ import annotations

import base64
import io

import pytest
from fastapi import HTTPException

from app.services.transcript import clean_vtt, parse_transcript

ZOOM_VTT = """WEBVTT

1
00:00:01.000 --> 00:00:04.000
Priya Sharma: I'll send the draft

2
00:00:04.000 --> 00:00:06.500
Priya Sharma: by Friday.

3
00:00:07.000 --> 00:00:09.000
Arjun Kumar: Agreed, let's ship Postgres 16.
"""

TEAMS_VTT = (
    "﻿WEBVTT\r\n\r\n"
    "6f1c2a9e-0b1d-4c55-9a1e-3d2f8f7c1a00/17-0\r\n"
    "00:00:01.000 --> 00:00:03.000\r\n"
    "<v Priya Sharma>We should cache the results.</v>\r\n\r\n"
    "NOTE this block is a comment\r\n\r\n"
    "6f1c2a9e-0b1d-4c55-9a1e-3d2f8f7c1a00/18-0\r\n"
    "00:00:03.000 --> 00:00:05.000\r\n"
    "<v Arjun Kumar>Fine by me,\r\n"
    "starting tomorrow.</v>\r\n"
)


def test_zoom_vtt_drops_timing_and_merges_a_speakers_cues():
    assert clean_vtt(ZOOM_VTT) == (
        "Priya Sharma: I'll send the draft by Friday.\n"
        "Arjun Kumar: Agreed, let's ship Postgres 16."
    )


def test_teams_vtt_with_bom_crlf_voice_tags_and_uuid_cue_ids():
    assert clean_vtt(TEAMS_VTT) == (
        "Priya Sharma: We should cache the results.\n"
        "Arjun Kumar: Fine by me, starting tomorrow."
    )


def test_vtt_upload_is_cleaned():
    assert parse_transcript("meeting.vtt", ZOOM_VTT.encode()).startswith("Priya Sharma: I'll send the draft")


def test_docx_transcript_is_read():
    import docx

    document = docx.Document()
    document.add_paragraph("Priya Sharma: I'll send the draft by Friday.")
    document.add_paragraph("   ")
    document.add_paragraph("Arjun Kumar: Agreed.")
    buf = io.BytesIO()
    document.save(buf)
    assert parse_transcript("Meeting notes.DOCX", buf.getvalue()) == (
        "Priya Sharma: I'll send the draft by Friday.\nArjun Kumar: Agreed."
    )


def test_windows_1252_text_is_decoded():
    assert parse_transcript("notes.txt", "Café sync: ship it\r\n".encode("cp1252")) == "Café sync: ship it"


def test_unsupported_type_and_empty_files_are_rejected():
    with pytest.raises(ValueError, match=r"\.vtt, \.txt or \.docx"):
        parse_transcript("slides.pdf", b"%PDF")
    with pytest.raises(ValueError, match="empty"):
        parse_transcript("notes.txt", b"  \r\n ")
    with pytest.raises(ValueError, match="could not be read"):
        parse_transcript("broken.docx", b"not a zip")


def test_parse_endpoint():
    from app.api.endpoints import TranscriptParseRequest, parse_meeting_transcript

    ok = parse_meeting_transcript(TranscriptParseRequest(filename="m.vtt", content_base64=base64.b64encode(ZOOM_VTT.encode()).decode()))
    assert ok["data"]["text"].startswith("Priya Sharma:")
    with pytest.raises(HTTPException) as bad:
        parse_meeting_transcript(TranscriptParseRequest(filename="m.vtt", content_base64="***"))
    assert bad.value.status_code == 400
    with pytest.raises(HTTPException) as wrong:
        parse_meeting_transcript(TranscriptParseRequest(filename="m.pdf", content_base64=base64.b64encode(b"x").decode()))
    assert wrong.value.status_code == 400
