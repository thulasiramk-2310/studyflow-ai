"""Every file type the upload accepts must actually yield text for indexing."""

from __future__ import annotations

import pytest

from app.loaders.document_loader import extract_text_from_document


def _text(pages):
    return "\n".join(p["text"] for p in pages)


def test_docx_paragraphs_and_tables_are_extracted(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("Deadlocks", level=1)
    d.add_paragraph("A deadlock needs mutual exclusion, hold and wait, no preemption and circular wait.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Banker's algorithm"
    table.rows[0].cells[1].text = "keeps the system in a safe state"
    path = tmp_path / "notes.docx"
    d.save(path)

    pages = extract_text_from_document(str(path))

    text = _text(pages)
    assert "circular wait" in text and "safe state" in text
    assert all(p["page_num"] >= 1 for p in pages)


def test_pptx_text_is_extracted_one_page_per_slide(tmp_path):
    from pptx import Presentation

    deck = Presentation()
    for title, body in (("Scheduling", "Round Robin uses a time quantum."), ("Paging", "Pages map to frames.")):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    path = tmp_path / "slides.pptx"
    deck.save(path)

    pages = extract_text_from_document(str(path))

    assert [p["page_num"] for p in pages] == [1, 2]
    assert "time quantum" in pages[0]["text"] and "frames" in pages[1]["text"]


@pytest.mark.parametrize("name", ["notes.md", "notes.txt"])
def test_markdown_and_plain_text_are_read(tmp_path, name):
    path = tmp_path / name
    path.write_text("# Paging\n\nPaging maps logical pages to physical frames.\n", encoding="utf-8")

    pages = extract_text_from_document(str(path))

    assert "physical frames" in _text(pages)


def test_unknown_types_are_still_rejected(tmp_path):
    path = tmp_path / "program.exe"
    path.write_bytes(b"MZ")
    with pytest.raises(ValueError):
        extract_text_from_document(str(path))


def test_indexing_keeps_the_uploaded_files_extension(monkeypatch):
    """The download used to be saved as .pdf whatever it was, so Word files never parsed."""
    from app.services import indexing

    seen = {}
    monkeypatch.setattr(indexing.storage, "download_to_file", lambda key, path: None)
    monkeypatch.setattr(indexing, "extract_text_from_document", lambda path: seen.setdefault("path", path) and [])
    monkeypatch.setattr(indexing, "update_resource_status", lambda *a, **k: None)

    indexing.process_document(resource_id=1, group_id=1, file_path="groups/1/uuid-1.docx", filename="notes.docx")

    assert seen["path"].endswith(".docx")
