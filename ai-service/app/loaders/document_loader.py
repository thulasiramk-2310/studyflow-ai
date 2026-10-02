import logging
from pathlib import Path

import fitz  # PyMuPDF

logger = logging.getLogger(__name__)

# Word documents have no fixed pages; group paragraphs into page-sized blocks so
# citations still point somewhere useful ("page 3" is roughly the third screenful).
DOCX_CHARS_PER_PAGE = 3000

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".pptx", ".md", ".txt")


def extract_text_from_document(file_path: str) -> list[dict]:
    """
    Extract text from a PDF, Word (.docx), PowerPoint (.pptx), Markdown or plain-text file.
    Returns a list of dictionaries: [{"page_num": int, "text": str}].
    Scanned PDFs (images without a text layer) yield little or no text: there is no OCR.
    """
    logger.info(f"Extracting text from {file_path}")
    ext = Path(file_path).suffix.lower()
    try:
        if ext == ".pdf":
            return _pdf(file_path)
        if ext == ".docx":
            return _docx(file_path)
        if ext == ".pptx":
            return _pptx(file_path)
        if ext in (".md", ".txt"):
            return [{"page_num": 1, "text": Path(file_path).read_text(encoding="utf-8", errors="replace")}]
        raise ValueError(f"Unsupported file type for {file_path}")
    except Exception as e:
        logger.error(f"Error extracting text from {file_path}: {e}")
        raise


def _pdf(file_path: str) -> list[dict]:
    pages = []
    doc = fitz.open(file_path)
    try:
        for page_num in range(len(doc)):
            pages.append({"page_num": page_num + 1, "text": doc.load_page(page_num).get_text()})
    finally:
        doc.close()
    return pages


def _docx(file_path: str) -> list[dict]:
    import docx

    document = docx.Document(file_path)
    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    pages, current = [], ""
    for block in blocks:
        if current and len(current) + len(block) > DOCX_CHARS_PER_PAGE:
            pages.append(current)
            current = ""
        current = f"{current}\n{block}" if current else block
    if current:
        pages.append(current)
    return [{"page_num": i + 1, "text": text} for i, text in enumerate(pages)]


def _pptx(file_path: str) -> list[dict]:
    from pptx import Presentation

    pages = []
    for number, slide in enumerate(Presentation(file_path).slides, start=1):
        parts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text.strip()]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append(slide.notes_slide.notes_text_frame.text)
        pages.append({"page_num": number, "text": "\n".join(parts)})
    return pages
